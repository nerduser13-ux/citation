"""Citation detection: patterns, char map across runs, exclusions, ambiguity."""

from __future__ import annotations

import pytest

from zoteroconv.detect_citations import (
    DOCUMENT_PART,
    FOOTNOTES_PART,
    detect_citations,
    expand_numbers,
)
from zoteroconv.docx_io import DocxPackage
from zoteroconv.parse_reference_list import find_reference_list


def _detect(path):
    pkg = DocxPackage.open(path)
    rl = find_reference_list(pkg)
    res = detect_citations(pkg, rl)
    # span consistency for every match, always
    for c in res.all_matches():
        pt = res.paragraph_texts[c.part][c.par_index]
        assert pt.text[c.start : c.end] == c.text
    return res


def _ok(res):
    return [(c.par_index, c.text, c.numbers, c.kind) for c in res.citations]


def test_simple_bracket_and_paren(gen_fixtures):
    res = _detect(gen_fixtures["simple_bracket.docx"])
    assert _ok(res) == [
        (0, "[1]", [1], "bracket"),
        (1, "(2)", [2], "paren"),
    ]
    assert not res.ambiguous and not res.excluded


def test_grouped_is_one_match(gen_fixtures):
    res = _detect(gen_fixtures["grouped_bracket.docx"])
    assert _ok(res) == [
        (0, "[2,5]", [2, 5], "bracket"),
        (1, "[1, 3-5]", [1, 3, 4, 5], "bracket"),
    ]


def test_ranges_with_all_dash_variants(gen_fixtures):
    res = _detect(gen_fixtures["range_dash.docx"])
    assert _ok(res) == [
        (0, "[3-6]", [3, 4, 5, 6], "bracket"),
        (1, "[3\u20136]", [3, 4, 5, 6], "bracket"),
        (2, "[3\u20146]", [3, 4, 5, 6], "bracket"),
    ]


def test_superscript_with_decoys(gen_fixtures):
    res = _detect(gen_fixtures["superscript.docx"])
    assert _ok(res) == [
        (0, "1,2", [1, 2], "superscript"),
        (0, "3", [3], "superscript"),
    ]
    reasons = [(c.par_index, c.text, c.reason) for c in res.ambiguous]
    assert (1, "1", "ordinal-like superscript (letter follows digits)") in reasons
    assert any(
        p == 1 and t == "2" and "unit token" in r for p, t, r in reasons
    )
    # every superscript char of an ok match really is superscript
    for c in res.citations:
        pt = res.paragraph_texts[c.part][c.par_index]
        assert all(ch.superscript for ch in pt.chars[c.start : c.end])


def test_citation_split_over_runs(gen_fixtures):
    res = _detect(gen_fixtures["split_runs.docx"])
    assert _ok(res) == [
        (0, "[2,5]", [2, 5], "bracket"),
        (1, "(1)", [1], "paren"),
    ]
    # the [2,5] match spans THREE different w:t elements (char map proof)
    c = res.citations[0]
    pt = res.paragraph_texts[c.part][c.par_index]
    t_elems = {id(ch.t_elem) for ch in pt.chars[c.start : c.end]}
    assert len(t_elems) == 3
    # the (1) match spans two w:t elements
    c2 = res.citations[1]
    pt2 = res.paragraph_texts[c2.part][c2.par_index]
    assert len({id(ch.t_elem) for ch in pt2.chars[c2.start : c2.end]}) == 2


def test_two_citations_in_one_paragraph(gen_fixtures):
    res = _detect(gen_fixtures["two_in_paragraph.docx"])
    assert _ok(res) == [
        (0, "[1]", [1], "bracket"),
        (0, "[2,3]", [2, 3], "bracket"),
    ]
    assert res.citations[0].start < res.citations[1].start


def test_citation_in_footnote(gen_fixtures):
    res = _detect(gen_fixtures["footnote_citation.docx"])
    parts = {(c.part, c.text, tuple(c.numbers)) for c in res.citations}
    assert (DOCUMENT_PART, "[1]", (1,)) in parts
    assert (FOOTNOTES_PART, "[2]", (2,)) in parts
    assert len(res.citations) == 2


def test_numbers_in_table_rules(gen_fixtures):
    res = _detect(gen_fixtures["numbers_in_table.docx"])
    # bracket form inside a table converts; bare/paren/superscript do not
    assert any(c.text == "[1]" and c.numbers == [1] for c in res.citations)
    assert any(c.text == "[2]" and c.numbers == [2] for c in res.citations)
    excluded = [(c.text, c.reason) for c in res.excluded]
    assert ("(2)", "number in a table (only bracket form converts there)") in excluded
    assert ("3", "number in a table (only bracket form converts there)") in excluded
    # the bare "5" in the table matches nothing at all
    assert all(c.text != "5" for c in res.all_matches())


def test_table_label_exclusion(gen_fixtures):
    res = _detect(gen_fixtures["table_label_exclusion.docx"])
    assert _ok(res) == [(2, "[2]", [2], "bracket")]
    amb = [(c.par_index, c.text, c.reason) for c in res.ambiguous]
    assert (0, "[2]", "follows a Table/Figure/Eq./Section/... label") in amb
    assert (1, "(1)", "follows a Table/Figure/Eq./Section/... label") in amb


def test_missing_reference_number_is_ambiguous(gen_fixtures):
    res = _detect(gen_fixtures["missing_ref_number.docx"])
    assert _ok(res) == [(0, "[4]", [4], "bracket")]
    assert len(res.ambiguous) == 1
    a = res.ambiguous[0]
    assert a.text == "[3]"
    assert "not present in the reference list" in a.reason


def test_reference_list_itself_never_scanned(gen_fixtures):
    res = _detect(gen_fixtures["reflist_not_converted.docx"])
    assert _ok(res) == [(0, "[1]", [1], "bracket")]
    assert not res.ambiguous and not res.excluded
    # no match may live at/after the References heading paragraph
    # (same package instance, so element identity holds)
    pkg = DocxPackage.open(gen_fixtures["reflist_not_converted.docx"])
    rl = find_reference_list(pkg)
    res2 = detect_citations(pkg, rl)
    heading_pts = [
        (part, pt.index)
        for part, pts in res2.paragraph_texts.items()
        for pt in pts
        if pt.par is rl.heading
    ]
    assert heading_pts
    part, hidx = heading_pts[0]
    assert all(
        c.part != part or c.par_index < hidx for c in res2.all_matches()
    )


def test_appendix_content_never_scanned(gen_fixtures):
    res = _detect(gen_fixtures["appendix_after_refs.docx"])
    assert _ok(res) == [(0, "[1]", [1], "bracket")]


def test_auto_numbered_refs_citations_valid(gen_fixtures):
    res = _detect(gen_fixtures["auto_numbered_refs.docx"])
    assert _ok(res) == [
        (3, "[2]", [2], "bracket"),
        (3, "[4]", [4], "bracket"),
    ]
    assert res.ref_numbers == {1, 2, 3, 4}


def test_partly_converted_decisions(gen_fixtures):
    res = _detect(gen_fixtures["partly_converted.docx"])
    assert _ok(res) == [
        (2, "[1]", [1], "bracket"),
        (2, "[2,3]", [2, 3], "bracket"),
        (3, "[4]", [4], "bracket"),      # raw [4] adjacent to existing field (decision 4)
        (4, "[1,4]", [1, 4], "bracket"),
    ]
    assert not res.ambiguous and not res.excluded
    # three existing fields were found and skipped (decisions 1+2)
    assert len(res.existing_fields) == 3
    assert any("skipped" in n for n in res.notes)
    # nothing matched inside a field zone
    for c in res.all_matches():
        pt = res.paragraph_texts[c.part][c.par_index]
        assert not any(ch.field_zone for ch in pt.chars[c.start : c.end])


def test_real_vancouver_twelve_citations(real_docx):
    res = _detect(real_docx)
    assert _ok(res) == [
        (2, "[8]", [8], "bracket"),
        (9, "[15]", [15], "bracket"),
        (32, "[9,11]", [9, 11], "bracket"),
        (33, "[4]", [4], "bracket"),
        (33, "[12]", [12], "bracket"),
        (34, "[10,13]", [10, 13], "bracket"),
        (37, "[7]", [7], "bracket"),
        (39, "[7]", [7], "bracket"),
        (39, "[16]", [16], "bracket"),
        (44, "[5]", [5], "bracket"),
        (44, "[14]", [14], "bracket"),
        (45, "[6]", [6], "bracket"),
    ]
    assert not res.ambiguous
    assert not res.excluded
    assert res.ref_numbers == set(range(1, 17))
    assert len(res.existing_fields) == 2


def test_sample_has_nothing_to_convert(sample_docx):
    res = _detect(sample_docx)
    assert res.citations == []
    assert res.ambiguous == []
    assert res.excluded == []


def test_enumeration_and_context_guards(tmp_path):
    from docx import Document

    p = tmp_path / "enumeration.docx"
    doc = Document()
    doc.add_paragraph("(2) the second option is worse.")          # start-of-paragraph
    doc.add_paragraph("Options: (1) first and (3) third.")        # after colon / lowercase
    doc.add_paragraph("As shown elsewhere(3) the effect holds.")  # glued to word -> cite
    doc.add_paragraph("Claims follow prior work (1), as expected.")  # , follows -> cite
    doc.add_paragraph("References")
    doc.add_paragraph("1. A A. Enum one. Journal E. 2020;1(1):1-5.")
    doc.add_paragraph("2. B B. Enum two. Journal F. 2021;2(2):6-9.")
    doc.add_paragraph("3. C C. Enum three. Journal G. 2022;3(3):10-15.")
    doc.save(p)

    res = _detect(p)
    assert _ok(res) == [
        (2, "(3)", [3], "paren"),
        (3, "(1)", [1], "paren"),
    ]
    amb = [(c.par_index, c.text, c.reason) for c in res.ambiguous]
    assert (0, "(2)", "enumeration-like: (n) at paragraph start") in amb
    assert (1, "(1)", "enumeration-like: (n) directly after a colon") in amb
    assert (1, "(3)", "enumeration-like: (n) followed by ' lowercase-word'") in amb


def test_expand_numbers_edge_cases():
    assert expand_numbers("2,5") == ([2, 5], None)
    assert expand_numbers("3-6") == ([3, 4, 5, 6], None)
    assert expand_numbers("1, 3\u20135") == ([1, 3, 4, 5], None)
    nums, err = expand_numbers("5-3")
    assert nums is None and "descending" in err
    nums, err = expand_numbers("1-999")
    assert nums is None and "too wide" in err


def test_dosage_unit_is_ambiguous(tmp_path):
    from docx import Document

    p = tmp_path / "dosage.docx"
    doc = Document()
    doc.add_paragraph("The dose was reduced [1] mg per day.")
    doc.add_paragraph("Prior evidence supports [1] strongly.")
    doc.add_paragraph("References")
    doc.add_paragraph("1. D D. Dose study. Journal D. 2020;1(1):1-5.")
    doc.save(p)
    res = _detect(p)
    assert _ok(res) == [(1, "[1]", [1], "bracket")]
    assert len(res.ambiguous) == 1
    assert "measurement unit" in res.ambiguous[0].reason
