"""
Phase 4 tests for the Typer CLI: usage errors, credential hygiene,
end-to-end conversion with validation, the loud mock-key failure, dry
run, --report, --limit-paragraphs excerpt mode.  HTTP is mocked through
the single zoteroconv.cli._make_transport seam.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from typer.testing import CliRunner

from test_convert import lib3, lib5
from test_resolve_items import (SECRET_KEY, FakeTransport, crossref_route,
                                dead_local, local_routes, web_routes)

from zoteroconv.cli import app
from zoteroconv.docx_io import DocxPackage

runner = CliRunner()
ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def patch_transport(monkeypatch):
    def _patch(routes):
        t = FakeTransport(routes)
        monkeypatch.setattr("zoteroconv.cli._make_transport", lambda: t)
        return t
    return _patch


def invoke(*args, **kw):
    return runner.invoke(app, [str(a) for a in args], **kw)


BASE = ["--style", "elsevier-harvard", "--no-cache"]


# ---------------------------------------------------------------------------
# usage / option validation
# ---------------------------------------------------------------------------

def test_cli_help():
    r = runner.invoke(app, ["--help"])
    assert r.exit_code == 0
    for flag in ("--style", "--output", "--dry-run", "--report",
                 "--limit-paragraphs", "--keep-uncited", "--zotero-read",
                 "--max-new-items", "--locale", "--collection",
                 "--source-style", "--api-key", "--library-id",
                 "--library-type"):
        assert flag in r.output


def test_cli_bad_style_lists_valid_options(gen_fixtures):
    # 'not-a-style' would PASS as an arbitrary CSL slug (by design); a
    # style with spaces/punctuation is clearly unusable
    r = invoke(gen_fixtures["simple_bracket.docx"], "--style",
               "Not A Style!", "--dry-run")
    assert r.exit_code == 2
    msg = r.output + r.stderr
    assert "apa" in msg and "harvard" in msg and "vancouver" in msg
    assert "zotero.org/styles" in msg
    r2 = invoke(gen_fixtures["simple_bracket.docx"], "--style",
                "https://example.com/styles/apa", "--dry-run")
    assert r2.exit_code == 2
    assert "not a zotero.org style URL" in (r2.output + r2.stderr)


def test_cli_source_style_only_auto(gen_fixtures):
    r = invoke(gen_fixtures["simple_bracket.docx"], "--style", "apa",
               "--source-style", "numeric", "--dry-run")
    assert r.exit_code == 2
    assert "only 'auto'" in (r.output + r.stderr)


def test_cli_output_required_unless_dry_run(gen_fixtures):
    r = invoke(gen_fixtures["simple_bracket.docx"], "--style", "apa")
    assert r.exit_code == 2
    assert "--output is required" in (r.output + r.stderr)


def test_cli_refuses_input_as_output(gen_fixtures):
    src = gen_fixtures["simple_bracket.docx"]
    r = invoke(src, "--style", "apa", "--output", src)
    assert r.exit_code == 2
    assert "refusing" in (r.output + r.stderr)


def test_cli_missing_input():
    r = invoke("/nonexistent/nope.docx", "--style", "apa", "--dry-run")
    assert r.exit_code != 0        # typer's exists=True check


# ---------------------------------------------------------------------------
# end-to-end with mocked HTTP
# ---------------------------------------------------------------------------

def test_cli_convert_validates_ok(gen_fixtures, tmp_path, patch_transport):
    patch_transport(local_routes(lib3()) + [crossref_route([])])
    out = tmp_path / "cli_out.docx"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--output", out,
               "--zotero-read", "local")
    assert r.exit_code == 0, r.output + r.stderr
    assert out.exists()
    assert "wrote 2 citation field(s)" in r.output
    assert "validation of" in r.output and ": OK" in r.output
    assert "library check: performed against the real Zotero library" \
        in r.output


def test_cli_exit1_when_validation_fails_on_mock_keys(
        gen_fixtures, tmp_path, patch_transport):
    # resolve+convert against the fake WEB API (creates NEW* keys); the
    # validation re-read of the library returns nothing -> loud failure
    patch_transport(dead_local() + web_routes(items=[]) + [crossref_route([])])
    out = tmp_path / "mock_out.docx"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--output", out,
               "--api-key", SECRET_KEY, "--library-id", "12345")
    assert out.exists()                      # conversion itself succeeded
    assert r.exit_code == 1
    assert "FAILED" in r.output
    assert "NOT FOUND" in r.output and "do NOT ship" in r.output
    assert "library-key" in r.output


def test_cli_dry_run_writes_nothing(gen_fixtures, tmp_path, patch_transport):
    patch_transport(local_routes(lib3()) + [crossref_route([])])
    out = tmp_path / "never.docx"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--dry-run",
               "--zotero-read", "local")
    assert r.exit_code == 0, r.output + r.stderr
    assert not out.exists()
    assert "dry-run: no file written" in r.output
    assert "validation" not in r.output.lower() or "SKIPPED" not in r.output


def test_cli_report_and_credential_hygiene(gen_fixtures, tmp_path,
                                           patch_transport):
    patch_transport(dead_local() + web_routes(items=[]) + [crossref_route([])])
    out = tmp_path / "rep_out.docx"
    rep = tmp_path / "nested" / "report.md"
    secret = "SUPER-SECRET-KEY-9999"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--output", out,
               "--report", rep, "--skip-library-check",
               "--api-key", secret, "--library-id", "12345")
    assert r.exit_code == 0, r.output + r.stderr
    assert secret not in r.output and secret not in r.stderr
    md = rep.read_text(encoding="utf-8")
    assert secret not in md
    assert "12345" not in md
    assert "provided (never logged)" in md
    assert "Open in Word with the Zotero plugin and click Zotero > Refresh." \
        in md
    assert "UNVERIFIED" in md
    assert "http://www.zotero.org/styles/elsevier-harvard" in md
    assert "PROVISIONAL" in md


def test_cli_env_credentials_hygiene(gen_fixtures, tmp_path,
                                     patch_transport, monkeypatch):
    patch_transport(dead_local() + web_routes(items=[]) + [crossref_route([])])
    monkeypatch.setenv("ZT_API_KEY", "ENV-SECRET-777")
    monkeypatch.setenv("ZT_LIBRARY_ID", "12345")   # what web_routes() serves
    monkeypatch.setenv("ZT_LIBRARY_TYPE", "user")
    out = tmp_path / "env_out.docx"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--output", out,
               "--skip-library-check")
    assert r.exit_code == 0, r.output + r.stderr
    combined = r.output + r.stderr
    assert "ENV-SECRET-777" not in combined
    assert "WEB API" in combined                 # writes need web creds


def test_cli_limit_paragraphs_excerpt(gen_fixtures, tmp_path,
                                      patch_transport):
    # two_in_paragraph: both citations live in paragraph #0
    patch_transport(local_routes(lib3()) + [crossref_route([])])
    out1 = tmp_path / "lim1.docx"
    r = invoke(gen_fixtures["two_in_paragraph.docx"], *BASE, "--output", out1,
               "--limit-paragraphs", "1", "--zotero-read", "local")
    assert r.exit_code == 0, r.output + r.stderr
    assert "excerpt mode" in r.output
    assert "wrote 2 citation field(s)" in r.output

    # limit 0 -> nothing converted, document intact, validation says so
    out0 = tmp_path / "lim0.docx"
    r0 = invoke(gen_fixtures["two_in_paragraph.docx"], *BASE, "--output", out0,
                "--limit-paragraphs", "0", "--zotero-read", "local",
                "--no-validate")
    assert r0.exit_code == 0, r0.output + r0.stderr
    assert "wrote 0 citation field(s)" in r0.output
    assert "beyond --limit-paragraphs 0" in r0.output
    conv = DocxPackage.open(out0)
    text = "".join(t.text or "" for t in
                   conv.document.xpath("//w:t", namespaces={
                       "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))
    assert "[1]" in text and "[2,3]" in text
    assert "Title one of the study" in text      # ref list untouched


def test_cli_limit_paragraphs_splits_document(tmp_path, patch_transport):
    """A citation inside the limit converts; one outside stays raw text."""
    from docx import Document
    src = tmp_path / "excerpt_src.docx"
    doc = Document()
    doc.add_paragraph("Intro claim [1] here.")
    doc.add_paragraph("Filler paragraph one.")
    doc.add_paragraph("Later result [2] discussed.")
    doc.add_paragraph("References")
    doc.add_paragraph("1. Smith A, Jones B. Title one of the study. "
                      "Journal One. 2020;1(1):1-10. "
                      "https://doi.org/10.1000/j1.2020.1")
    doc.add_paragraph("2. Lee C. Title two of review. Journal Two. "
                      "2021;2(2):11-20. https://doi.org/10.1000/j2.2021.2")
    doc.save(src)
    patch_transport(local_routes(lib3()) + [crossref_route([])])
    out = tmp_path / "excerpt_out.docx"
    r = invoke(src, *BASE, "--output", out, "--limit-paragraphs", "2",
               "--zotero-read", "local", "--report", tmp_path / "exc.md")
    assert r.exit_code == 0, r.output + r.stderr
    assert "wrote 1 citation field(s)" in r.output
    assert "beyond --limit-paragraphs 2" in r.output
    md = (tmp_path / "exc.md").read_text(encoding="utf-8")
    assert "EXCERPT MODE" in md
    assert "[2]" in md and "excerpt mode" in md


def test_cli_keep_uncited_flag(gen_fixtures, tmp_path, patch_transport):
    patch_transport(local_routes(lib3()) + [crossref_route([])])
    out = tmp_path / "keep.docx"
    r = invoke(gen_fixtures["simple_bracket.docx"], *BASE, "--output", out,
               "--keep-uncited", "--zotero-read", "local")
    assert r.exit_code == 0, r.output + r.stderr
    assert "KEPT via uncited array" in r.output
    assert "UNVERIFIED" in r.output


def test_cli_grouped_and_superscript(gen_fixtures, tmp_path,
                                     patch_transport):
    patch_transport(local_routes(lib5()) + [crossref_route([])])
    out = tmp_path / "grouped.docx"
    r = invoke(gen_fixtures["grouped_bracket.docx"],
               "--style", "vancouver", "--no-cache", "--output", out,
               "--zotero-read", "local")
    assert r.exit_code == 0, r.output + r.stderr
    assert "wrote 2 citation field(s)" in r.output
    assert "nlm-citation-sequence" in r.output   # dependent-style fallback


# ---------------------------------------------------------------------------
# module entry points
# ---------------------------------------------------------------------------

def test_python_dash_m_help():
    p = subprocess.run([sys.executable, "-m", "zoteroconv", "--help"],
                       capture_output=True, text=True, cwd=str(ROOT),
                       timeout=120)
    assert p.returncode == 0, p.stderr
    assert "--limit-paragraphs" in p.stdout


def test_python_dash_m_validate(sample_docx):
    p = subprocess.run([sys.executable, "-m", "zoteroconv.validate",
                        str(sample_docx), "--skip-library-check"],
                       capture_output=True, text=True, cwd=str(ROOT),
                       timeout=120)
    assert p.returncode == 0, p.stdout + p.stderr
    assert "OK" in p.stdout
