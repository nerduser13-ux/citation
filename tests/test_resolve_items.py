"""Phase 2: Zotero access + resolution.  ALL HTTP is mocked (FakeTransport);
the single @pytest.mark.live test is skipped unless --run-live is passed."""

from __future__ import annotations

import re

import pytest

from zoteroconv.docx_io import DocxPackage
from zoteroconv.http_cache import HttpCache
from zoteroconv.resolve_items import (
    CROSSREF_WORKS_URL,
    LOCAL_BASE,
    PUBMED_ESUMMARY_URL,
    WEB_BASE,
    Candidate,
    CrossRefClient,
    HttpResponse,
    LibraryItem,
    PubmedClient,
    RequestsTransport,
    Transport,
    ZoteroAccess,
    csl_to_zotero_item,
    evaluate_suggestion,
    library_items_from_csl,
    match_reference,
    resolve_document,
    stable_numeric_id,
)

LOCAL_ITEMS_URL = LOCAL_BASE + "/items"
LOCAL_COLL_URL = LOCAL_BASE + "/collections"
WEB = WEB_BASE + "/users/12345"
WEB_ITEMS_URL = WEB + "/items"
WEB_COLL_URL = WEB + "/collections"
SECRET_KEY = "ZT-SECRET-abc123"


# ---------------------------------------------------------------------------
# fake transport
# ---------------------------------------------------------------------------

class FakeTransport(Transport):
    def __init__(self, routes):
        self.routes = routes          # [(method, url_regex, handler|response)]
        self.calls: list[dict] = []

    def request(self, method, url, *, params=None, json=None, headers=None,
                timeout=10.0):
        call = {"method": method, "url": url, "params": params or {},
                "json": json, "headers": headers or {}}
        self.calls.append(call)
        for m, pat, handler in self.routes:
            if m == method and re.search(pat, url):
                out = handler(call) if callable(handler) else handler
                if isinstance(out, Exception):
                    raise out
                return out
        raise AssertionError(f"FakeTransport: unexpected HTTP {method} {url}")

    def calls_to(self, url_substring):
        return [c for c in self.calls if url_substring in c["url"]]

    def posts_to(self, url_substring):
        return [c for c in self.calls_to(url_substring) if c["method"] == "POST"]


def local_routes(items, *, write_status=405):
    def get_items(call):
        if call["params"].get("limit") == 1:      # reachability probe
            return HttpResponse(200, json=items[:1])
        return HttpResponse(200, json=items)

    def post_items(call):
        if call["json"] == [{}]:                  # write probe
            return HttpResponse(write_status, json=None, text="")
        body = call["json"]
        ok = {str(i): {"key": f"NEWL{i + 1}"} for i in range(len(body))}
        return HttpResponse(200, json={"success": ok, "unchanged": {},
                                       "failed": {}})

    return [
        ("GET", re.escape(LOCAL_ITEMS_URL), get_items),
        ("POST", re.escape(LOCAL_ITEMS_URL), post_items),
        ("GET", re.escape(LOCAL_COLL_URL), HttpResponse(200, json=[])),
        ("POST", re.escape(LOCAL_COLL_URL),
         HttpResponse(201, json={"success": {"0": {"key": "COLLL"}}})),
    ]


def web_routes(items=None, collections=None, key_prefix="NEW"):
    state = {"n": 0}

    def get_items(call):
        return HttpResponse(200, json=list(items or []))

    def post_items(call):
        body = call["json"]
        ok = {}
        for i in range(len(body)):
            state["n"] += 1
            ok[str(i)] = {"key": f"{key_prefix}{state['n']}"}
        return HttpResponse(200, json={"success": ok, "unchanged": {},
                                       "failed": {}})

    return [
        ("GET", re.escape(WEB_ITEMS_URL), get_items),
        ("GET", re.escape(WEB_COLL_URL),
         HttpResponse(200, json=list(collections or []))),
        ("POST", re.escape(WEB_COLL_URL),
         HttpResponse(201, json={"success": {"0": {"key": "COLL9"}}})),
        ("POST", re.escape(WEB_ITEMS_URL), post_items),
    ]


def dead_local():
    return [
        ("GET", re.escape(LOCAL_BASE), ConnectionRefusedError()),
        ("POST", re.escape(LOCAL_BASE), ConnectionRefusedError()),
    ]


def crossref_route(works):
    return ("GET", re.escape(CROSSREF_WORKS_URL),
            HttpResponse(200, json={"message": {"items": works}}))


def pubmed_route(result):
    return ("GET", re.escape(PUBMED_ESUMMARY_URL),
            HttpResponse(200, json={"result": result}))


def lib_item(key, token="TOKENLOC", **csl):
    csl.setdefault("id", f"http://zotero.org/users/local/{token}/items/{key}")
    return csl


def build_doc(tmp_path, name, body_paras, ref_paras):
    from docx import Document

    p = tmp_path / name
    doc = Document()
    for t in body_paras:
        doc.add_paragraph(t)
    doc.add_paragraph("References")
    for t in ref_paras:
        doc.add_paragraph(t)
    doc.save(p)
    return DocxPackage.open(p)


# ---------------------------------------------------------------------------
# read-mode detection + privacy
# ---------------------------------------------------------------------------

def test_local_read_mode_line_and_no_web_calls(gen_fixtures, capsys):
    items = [
        lib_item("ITM2KEY", type="article-journal",
                 title="Title two about things", DOI="10.1000/j2.2021.2",
                 author=[{"family": "Smith", "given": "J"},
                         {"family": "Jones", "given": "AB"}],
                 issued={"date-parts": [["2021", 2, 2]]}),
        lib_item("ITM3KEY", type="article-journal",
                 title="Title three of analysis",
                 author=[{"family": "Brown", "given": "C"}],
                 issued={"date-parts": [["2022"]]}),
    ]
    t = FakeTransport(local_routes(items))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t)
    out = capsys.readouterr().out

    assert rep.mode == "local"
    assert rep.local_writes is False            # probe answered 405
    assert rep.write_channel is None            # no web credentials
    assert "LOCAL API" in out and "refused" in out      # decision 1: mode printed
    assert SECRET_KEY not in out

    r1, r2, r3 = (rep.by_ref(n) for n in (1, 2, 3))
    # decision 7: matched items carry key / sample-form URI / numeric id / CSL
    assert r2.status == "matched" and r2.method == "doi"
    assert r2.item.key == "ITM2KEY"
    assert r2.item.uri == "http://zotero.org/users/local/TOKENLOC/items/ITM2KEY"
    assert isinstance(r2.item.numeric_id, int) and 10_000_000 <= r2.item.numeric_id < 100_000_000
    assert r2.item.csl["DOI"] == "10.1000/j2.2021.2"
    assert r3.status == "matched" and r3.method == "fuzzy"
    assert r3.item.key == "ITM3KEY"
    # ref1 matched nothing -> planned, then blocked before any write (decision 2)
    assert r1.status == "create-blocked-no-writes"
    assert any("STOPPED BEFORE ANY WRITE" in n for n in rep.notes)
    assert "ZT_API_KEY" in out or any("ZT_API_KEY" in n for n in rep.notes)

    # reads asked for CSL-JSON and never touched the web or CrossRef
    fetches = [c for c in t.calls_to(LOCAL_ITEMS_URL)
               if c["method"] == "GET" and c["params"].get("limit") == 100]
    assert fetches and fetches[0]["headers"]["Accept"] == (
        "application/vnd.citationstyles.csl+json"
    )
    for c in t.calls_to("localhost"):
        assert "Zotero-API-Key" not in c["headers"]
    assert t.calls_to("api.zotero.org") == []
    assert t.calls_to("crossref.org") == []     # ref1 HAS a DOI: no suggestion


def test_web_mode_credentials_header_and_privacy(gen_fixtures, capsys):
    t = FakeTransport(dead_local() + web_routes(items=[]))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345", library_type="user")
    out = capsys.readouterr().out

    assert rep.mode == "web" and rep.write_channel == "web"
    assert "WEB API" in out
    assert SECRET_KEY not in out                # decision 9: never printed
    web_calls = t.calls_to("api.zotero.org")
    assert web_calls
    assert all(c["headers"].get("Zotero-API-Key") == SECRET_KEY for c in web_calls)
    assert all(c["headers"].get("Zotero-API-Version") == "3" for c in web_calls)


def test_forced_read_modes():
    items = [lib_item("K1", title="T", issued={"date-parts": [["2020"]]})]
    # --zotero-read web: local never contacted even though it would answer
    t = FakeTransport(web_routes(items=items))
    acc = ZoteroAccess(t, api_key="k", library_id="12345", zotero_read="web")
    acc.detect()
    got = acc.fetch_items_csl()
    assert acc.read_mode == "web" and len(got) == 1
    assert t.calls_to("localhost") == []

    # --zotero-read local while local is down: offline, no web fallback
    t2 = FakeTransport(dead_local() + web_routes(items=items))
    acc2 = ZoteroAccess(t2, api_key="k", library_id="12345", zotero_read="local")
    acc2.detect()
    got2 = acc2.fetch_items_csl()
    assert acc2.read_mode == "offline" and got2 == []
    assert any("offline" in n for n in acc2.notes)
    assert t2.calls_to("api.zotero.org") == []


def test_local_write_probe_accepted_enables_local_writes(gen_fixtures):
    items = [lib_item("ITM2KEY", title="Title two about things",
                      DOI="10.1000/j2.2021.2",
                      author=[{"family": "Smith", "given": "J"},
                              {"family": "Jones", "given": "AB"}],
                      issued={"date-parts": [["2021", 2, 2]]})]
    # probe answered 400 => POST endpoint validates => writes accepted
    t = FakeTransport(local_routes(items, write_status=400))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t)
    assert rep.local_writes is True
    assert rep.write_channel == "local"
    created = [r for r in rep.resolutions if r.status == "created"]
    assert len(created) == 2                     # refs 1 and 3
    # URIs reuse the sample-form local prefix learned from fetched items
    uris = {r.item.uri for r in created}
    assert all(u.startswith("http://zotero.org/users/local/TOKENLOC/items/")
               for u in uris)
    assert t.posts_to(LOCAL_ITEMS_URL)           # writes went to the LOCAL API
    assert t.calls_to("api.zotero.org") == []


def test_fetch_falls_back_to_format_csljson():
    csl = [lib_item("K1", title="T one", issued={"date-parts": [["2020"]]})]

    def get_items(call):
        if call["params"].get("limit") == 1:          # reachability probe
            return HttpResponse(200, json=[])
        if call["params"].get("format") == "csljson":
            return HttpResponse(200, json=csl)        # this Zotero version...
        return HttpResponse(400, text="bad Accept")   # ...rejects the header

    t = FakeTransport([("GET", re.escape(LOCAL_ITEMS_URL), get_items),
                       ("POST", re.escape(LOCAL_ITEMS_URL),
                        HttpResponse(405))])
    acc = ZoteroAccess(t)
    acc.detect()
    items = acc.fetch_items_csl()
    assert items == csl
    assert any("format=csljson" in n for n in acc.notes)


def test_fetch_converts_native_zotero_json():
    zjson = [{"key": "ABCD1234",
              "data": {"key": "ABCD1234", "itemType": "journalArticle",
                       "title": "Converted title",
                       "creators": [{"creatorType": "author",
                                     "lastName": "Smith", "firstName": "J"}],
                       "date": "2021-06-03",
                       "publicationTitle": "Journal of Conversion",
                       "volume": "9", "issue": "2", "pages": "1-9",
                       "DOI": "10.1/conv", "extra": "PMID: 555"}}]
    t = FakeTransport([("GET", re.escape(LOCAL_ITEMS_URL),
                        HttpResponse(200, json=zjson))])
    acc = ZoteroAccess(t)
    acc.local_reachable = True
    acc.read_mode = "local"
    items = acc.fetch_items_csl()
    assert len(items) == 1
    c = items[0]
    assert c["type"] == "article-journal"
    assert c["title"] == "Converted title"
    assert c["author"] == [{"family": "Smith", "given": "J"}]
    assert c["issued"] == {"date-parts": [["2021", 6, 3]]}
    assert c["container-title"] == "Journal of Conversion"
    assert c["volume"] == "9" and c["issue"] == "2" and c["page"] == "1-9"
    assert c["DOI"] == "10.1/conv" and c["PMID"] == "555"
    assert c["id"].endswith("/items/ABCD1234")
    assert any("converted to CSL-JSON locally" in n for n in acc.notes)


def test_uncited_refs_left_out_by_default_and_kept_on_request(gen_fixtures):
    # default: cited_numbers={1} -> refs 2,3 resolved read-only, not created
    t = FakeTransport(dead_local() + web_routes(items=[]))
    rep = resolve_document(DocxPackage.open(gen_fixtures["simple_bracket.docx"]),
                           transport=t, api_key=SECRET_KEY, library_id="12345",
                           cited_numbers={1})
    assert rep.uncited == [2, 3]
    assert rep.by_ref(1).status == "created"
    assert rep.by_ref(2).status == "uncited-not-created"
    assert rep.by_ref(3).status == "uncited-not-created"
    assert any("never cited numerically" in n for n in rep.notes)
    assert len(t.posts_to(WEB_ITEMS_URL)) == 1

    # --keep-uncited: they get items like everything else
    t2 = FakeTransport(dead_local() + web_routes(items=[]))
    rep2 = resolve_document(DocxPackage.open(gen_fixtures["simple_bracket.docx"]),
                            transport=t2, api_key=SECRET_KEY, library_id="12345",
                            cited_numbers={1}, keep_uncited=True)
    assert [rep2.by_ref(n).status for n in (1, 2, 3)] == ["created"] * 3


def test_web_pagination():
    page1 = [lib_item(f"K{i:03d}", title=f"T{i}") for i in range(100)]
    page2 = [lib_item(f"K{i:03d}", title=f"T{i}") for i in range(100, 103)]

    def get_items(call):
        start = call["params"].get("start", 0)
        return HttpResponse(200, json=page1 if start == 0 else page2)

    t = FakeTransport([("GET", re.escape(LOCAL_ITEMS_URL), get_items)])
    acc = ZoteroAccess(t)
    acc.local_reachable = True
    acc.read_mode = "local"
    items = acc.fetch_items_csl()
    assert len(items) == 103
    assert len(t.calls_to(LOCAL_ITEMS_URL)) == 2


# ---------------------------------------------------------------------------
# match cascade
# ---------------------------------------------------------------------------

def _lib(items):
    return library_items_from_csl(items, "local")


def test_cascade_doi_pmid_isbn_exact():
    lib = _lib([
        lib_item("A", title="X", DOI="10.1/Abc"),
        lib_item("B", title="Y", PMID="999"),
        lib_item("C", title="Z", ISBN="978-0-306-40615-7"),
    ])
    m = match_reference({"title": "whatever", "DOI": "https://doi.org/10.1/abc"}, lib)
    assert m.method == "doi" and m.item.key == "A"
    m = match_reference({"title": "whatever", "PMID": "999"}, lib)
    assert m.method == "pmid" and m.item.key == "B"
    m = match_reference({"title": "whatever", "ISBN": "9780306406157"}, lib)
    assert m.method == "isbn" and m.item.key == "C"


def test_cascade_fuzzy_single_and_multiple():
    lib = _lib([
        lib_item("A", title="Deep learning for financial services",
                 author=[{"family": "Smith", "given": "J"}],
                 issued={"date-parts": [["2021"]]}),
    ])
    m = match_reference(
        {"title": "Deep learning for financial services!",
         "author": [{"family": "Smith", "given": "J."}],
         "issued": {"date-parts": [["2021"]]}}, lib)
    assert m.method == "fuzzy" and m.item.key == "A"
    assert m.candidates[0].score >= 90

    lib2 = lib + _lib([
        lib_item("B", title="Deep learning for financial service",
                 author=[{"family": "Smith", "given": "J"}],
                 issued={"date-parts": [["2021"]]}),
    ])
    m2 = match_reference(
        {"title": "Deep learning for financial services",
         "author": [{"family": "Smith", "given": "J"}],
         "issued": {"date-parts": [["2021"]]}}, lib2)
    # decision 3: two candidates above threshold -> never auto-picked
    assert m2.item is None and m2.method is None
    assert len(m2.candidates) == 2
    assert any("not auto-picked" in msg for msg in m2.messages)


def test_cascade_no_title_no_match():
    m = match_reference({"author": [{"family": "Smith"}]}, _lib([lib_item("A", title="T")]))
    assert m.item is None
    assert any("no usable title" in msg for msg in m.messages)


def test_evaluate_suggestion_rules():
    ref = {"title": "Title one of the study",
           "author": [{"family": "Author", "given": "A"}],
           "issued": {"date-parts": [["2020"]]}}
    ok, sim, _ = evaluate_suggestion(ref, {
        "title": "Title one of the study", "year": "2020", "surname": "Author"})
    assert ok and sim > 95
    # year mismatch
    ok, _, why = evaluate_suggestion(ref, {
        "title": "Title one of the study", "year": "2019", "surname": "Author"})
    assert not ok and "MISMATCH" in why
    # author mismatch
    ok, _, why = evaluate_suggestion(ref, {
        "title": "Title one of the study", "year": "2020", "surname": "Zother"})
    assert not ok and "MISMATCH" in why
    # weak title
    ok, _, _ = evaluate_suggestion(ref, {
        "title": "Completely different work", "year": "2020", "surname": "Author"})
    assert not ok


def test_stable_numeric_id():
    a = stable_numeric_id("http://zotero.org/users/local/T/items/K1")
    b = stable_numeric_id("http://zotero.org/users/local/T/items/K1")
    c = stable_numeric_id("http://zotero.org/users/local/T/items/K2")
    assert a == b and a != c
    assert all(10_000_000 <= x < 100_000_000 for x in (a, b, c))


# ---------------------------------------------------------------------------
# CrossRef / PubMed suggestions
# ---------------------------------------------------------------------------

def test_crossref_accept_matches_library_doi(tmp_path):
    pkg = build_doc(
        tmp_path, "nodoc.docx",
        ["Intro cites [1]."],
        ["1. Author A, Author B. Title one of the study. Journal One. "
         "2020;1(1):1-10."],          # no DOI/PMID/ISBN
    )
    work = {"DOI": "10.1000/CR.2020.99",
            "title": ["Title one of the study"],
            "issued": {"date-parts": [[2020, 1, 15]]},
            "author": [{"family": "Author", "given": "A."}]}
    lib = [lib_item("ITMX", title="Completely different canonical title xyz",
                    DOI="10.1000/CR.2020.99", issued={"date-parts": [["2019"]]})]
    t = FakeTransport(local_routes(lib) + [crossref_route([work])])
    rep = resolve_document(pkg, transport=t, crossref_mailto="me@example.org")
    r1 = rep.by_ref(1)
    assert r1.status == "matched"
    assert r1.method == "crossref-doi->doi"
    assert r1.item.key == "ITMX"
    # decision 5: queried with the bibliographic string only + polite UA/mailto
    # (raw has the typed "1." prefix stripped by the parser - exactly the
    # bibliographic string)
    calls = t.calls_to("api.crossref.org")
    assert len(calls) == 1
    assert calls[0]["params"]["query.bibliographic"] == (
        "Author A, Author B. Title one of the study. Journal One. "
        "2020;1(1):1-10."
    )
    assert calls[0]["params"]["mailto"] == "me@example.org"
    assert "mailto:me@example.org" in calls[0]["headers"]["User-Agent"]


def test_crossref_weak_suggestion_keeps_placeholder(tmp_path):
    pkg = build_doc(
        tmp_path, "weakdoc.docx",
        ["Intro cites [1]."],
        ["1. Author A, Author B. Title one of the study. Journal One. "
         "2020;1(1):1-10."],
    )
    work = {"DOI": "10.9999/wrong.suggestion",
            "title": ["Unrelated other work entirely"],
            "issued": {"date-parts": [[2020]]},
            "author": [{"family": "Author", "given": "A."}]}
    t = FakeTransport(dead_local() + web_routes(items=[]) + [crossref_route([work])])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY, library_id="12345")
    r1 = rep.by_ref(1)
    assert r1.status == "created"                    # placeholder got created ...
    assert r1.accepted_doi is None                   # ... WITHOUT the suggestion
    assert r1.candidates and r1.candidates[0].source == "crossref"
    assert r1.candidates[0].doi == "10.9999/wrong.suggestion"
    assert any("no metadata taken" in m for m in r1.messages)
    body = t.posts_to(WEB_ITEMS_URL)[0]["json"][0]
    assert "DOI" not in body                         # never fabricated
    assert body["title"] == "Title one of the study"
    # user directive: low-confidence placeholders additionally get the review tag
    assert body["tags"] == [{"tag": "zoteroconv-created"},
                            {"tag": "zoteroconv-review"}]


def test_crossref_cache_hits_once(tmp_path):
    from zoteroconv.resolve_items import _sha1

    cache = HttpCache(tmp_path / "c.sqlite3")
    works = [{"DOI": "10.1/x", "title": ["X"]}]
    t = FakeTransport([crossref_route(works)])
    client = CrossRefClient(t, cache=cache, mailto=None)
    a = client.suggest("some bibliographic string")
    b = client.suggest("some bibliographic string")
    assert a == b == works and len(t.calls_to("api.crossref.org")) == 1
    key = _sha1("works|5|some bibliographic string")
    cache.close()
    with HttpCache(tmp_path / "c.sqlite3") as cache2:   # persisted
        assert cache2.get("crossref", key) == works


def test_pubmed_accept_and_create(tmp_path):
    pkg = build_doc(
        tmp_path, "pmdoc.docx",
        ["Intro cites [1]."],
        ["1. Med M. Medical study title here. Journal M. 2019;5(5):55-60. "
         "PMID: 12345678."],
    )
    result = {"12345678": {
        "title": "Medical study title here.",
        "pubdate": "2019 Jun 3",
        "authors": [{"name": "Med M"}],
        "articleids": [{"idtype": "doi", "value": "10.9999/med.2019"}],
    }}
    t = FakeTransport(dead_local() + web_routes(items=[]) + [pubmed_route(result)])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY, library_id="12345")
    r1 = rep.by_ref(1)
    assert r1.status == "created"
    assert r1.accepted_doi == "10.9999/med.2019"
    body = t.posts_to(WEB_ITEMS_URL)[0]["json"][0]
    assert body["DOI"] == "10.9999/med.2019"
    assert body["extra"] == "PMID: 12345678"
    pm_calls = t.calls_to("eutils.ncbi.nlm.nih.gov")
    assert pm_calls[0]["params"] == {"db": "pubmed", "id": "12345678",
                                     "retmode": "json"}


def test_pubmed_weak_suggestion_low_confidence(tmp_path):
    pkg = build_doc(
        tmp_path, "pmweak.docx",
        ["Intro cites [1]."],
        ["1. Med M. Medical study title here. Journal M. 2019;5(5):55-60. "
         "PMID: 12345678."],
    )
    result = {"12345678": {
        "title": "Some other medical paper.", "pubdate": "2019",
        "authors": [{"name": "Other P"}],
        "articleids": [{"idtype": "doi", "value": "10.1/other"}],
    }}
    t = FakeTransport(dead_local() + web_routes(items=[]) + [pubmed_route(result)])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY, library_id="12345")
    r1 = rep.by_ref(1)
    assert r1.status == "created"
    assert r1.accepted_doi is None
    assert any(c.source == "pubmed" for c in r1.candidates)
    body = t.posts_to(WEB_ITEMS_URL)[0]["json"][0]
    assert "DOI" not in body
    assert body["extra"] == "PMID: 12345678"
    assert {"tag": "zoteroconv-review"} in body["tags"]


def test_pubmed_cache(tmp_path):
    cache = HttpCache(tmp_path / "pm.sqlite3")
    t = FakeTransport([pubmed_route({"77": {"title": "T", "pubdate": "2020",
                                            "authors": [], "articleids": []}})])
    client = PubmedClient(t, cache=cache)
    client.suggest_for_pmid("77")
    client.suggest_for_pmid("77")
    assert len(t.calls_to("eutils")) == 1
    assert client.suggest_for_pmid("not-a-number") is None


# ---------------------------------------------------------------------------
# creates: collection, dry-run, cap+confirm, dedup, existing fields
# ---------------------------------------------------------------------------

def test_creates_with_collection_and_payload(gen_fixtures):
    t = FakeTransport(dead_local() + web_routes(items=[]))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345")
    assert rep.collection_name == "Imported - simple_bracket"
    assert rep.collection_key == "COLL9"
    post_coll = t.posts_to(WEB_COLL_URL)
    assert post_coll and post_coll[0]["json"] == [
        {"name": "Imported - simple_bracket"}]
    created = [r for r in rep.resolutions if r.status == "created"]
    assert [r.ref_number for r in created] == [1, 2, 3]
    assert rep.created_keys == ["NEW1", "NEW2", "NEW3"]

    r2 = rep.by_ref(2)
    assert r2.item.uri == "http://zotero.org/users/12345/items/NEW2"
    assert r2.item.csl["id"] == r2.item.uri
    assert 10_000_000 <= r2.item.numeric_id < 100_000_000
    body = t.posts_to(WEB_ITEMS_URL)[1]["json"][0]
    assert body["itemType"] == "journalArticle"
    assert body["title"] == "Title two about things"
    assert body["creators"] == [
        {"creatorType": "author", "lastName": "Smith", "firstName": "J"},
        {"creatorType": "author", "lastName": "Jones", "firstName": "AB"},
    ]
    assert body["date"] == "2021"
    assert body["publicationTitle"] == "Journal Two"
    assert body["volume"] == "2" and body["issue"] == "2"
    assert body["pages"] == "11-20"
    assert body["DOI"] == "10.1000/j2.2021.2"
    assert body["collections"] == ["COLL9"]      # decision 6: added to collection
    # user directive: every created item is tagged for cleanup
    assert body["tags"] == [{"tag": "zoteroconv-created"}]


def test_existing_collection_reused(gen_fixtures):
    colls = [{"key": "COLLX",
              "data": {"name": "Imported - simple_bracket"}}]
    t = FakeTransport(dead_local() + web_routes(items=[], collections=colls))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345")
    assert rep.collection_key == "COLLX"
    assert t.posts_to(WEB_COLL_URL) == []        # checked first, not re-created
    assert any("already exists" in n for n in rep.notes)


def test_dry_run_never_writes(gen_fixtures):
    t = FakeTransport(dead_local() + web_routes(items=[]))
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345", dry_run=True)
    assert rep.dry_run is True
    assert [r.status for r in rep.resolutions] == ["create-planned"] * 3
    assert all(c["method"] == "GET" for c in t.calls)   # decision 2: no writes
    assert any("dry-run" in n for n in rep.notes)


def test_max_new_items_cap_and_confirm(gen_fixtures):
    pkg = DocxPackage.open(gen_fixtures["grouped_bracket.docx"])  # 5 refs
    asked = []

    def confirm(q):
        asked.append(q)
        return False

    t = FakeTransport(dead_local() + web_routes(items=[]))
    rep = resolve_document(DocxPackage.open(gen_fixtures["grouped_bracket.docx"]),
                           transport=t, api_key=SECRET_KEY, library_id="12345",
                           max_new_items=2, confirm=confirm)
    assert len(asked) == 1 and "5 new items" in asked[0]
    assert "--max-new-items=2" in asked[0]
    counts = rep.counts()
    assert counts.get("created") == 2
    assert counts.get("skipped-max-new-items") == 3

    # confirm=True -> create all 5
    t2 = FakeTransport(dead_local() + web_routes(items=[]))
    rep2 = resolve_document(pkg, transport=t2, api_key=SECRET_KEY,
                            library_id="12345", max_new_items=2,
                            confirm=lambda q: True)
    assert rep2.counts().get("created") == 5


def test_duplicate_before_create_reuses_created_item(tmp_path):
    dup_ref = ("Dup D. Duplicate study of things. Journal X. 2020;1(1):1-5. "
               "https://doi.org/10.5555/dup.2020")
    pkg = build_doc(
        tmp_path, "dupdoc.docx",
        ["Intro cites [1] and [2]."],
        [f"1. {dup_ref}", f"2. {dup_ref}"],
    )
    t = FakeTransport(dead_local() + web_routes(items=[]))
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345")
    r1, r2 = rep.by_ref(1), rep.by_ref(2)
    assert r1.status == "created"
    assert r2.status == "duplicate-skip"
    assert r2.item.key == r1.item.key
    assert r2.method == "dedup-before-create:doi"
    assert len(t.posts_to(WEB_ITEMS_URL)) == 1   # only ONE item created


def test_existing_fields_reused_not_recreated(gen_fixtures):
    t = FakeTransport(local_routes([]))          # local read, writes refused
    pkg = DocxPackage.open(gen_fixtures["partly_converted.docx"])
    rep = resolve_document(pkg, transport=t)
    r4 = rep.by_ref(4)
    assert r4.status == "reused-existing-field"
    assert r4.method == "existing-field-doi"
    assert r4.item.key == "FIXITEM1"
    assert r4.item.uri == "http://zotero.org/users/local/FIXLIB01/items/FIXITEM1"
    assert r4.item.csl["DOI"] == "10.46299/j.isjmef.20240303.01"
    assert isinstance(r4.item.numeric_id, int)
    assert any("will not be created again" in m for m in r4.messages)
    # refs 1-3 blocked (no credentials); ref 4 never entered the create plan
    assert sorted(r.ref_number for r in rep.resolutions
                  if r.status == "create-blocked-no-writes") == [1, 2, 3]
    assert t.calls_to("api.crossref.org") == []


def test_create_failure_reported(gen_fixtures):
    def post_items(call):
        return HttpResponse(200, json={
            "success": {}, "unchanged": {},
            "failed": {"0": {"message": "Invalid item type"}},
        })

    routes = dead_local() + [
        ("GET", re.escape(WEB_ITEMS_URL), HttpResponse(200, json=[])),
        ("GET", re.escape(WEB_COLL_URL), HttpResponse(200, json=[])),
        ("POST", re.escape(WEB_COLL_URL),
         HttpResponse(201, json={"success": {"0": {"key": "C1"}}})),
        ("POST", re.escape(WEB_ITEMS_URL), post_items),
    ]
    t = FakeTransport(routes)
    pkg = DocxPackage.open(gen_fixtures["simple_bracket.docx"])
    rep = resolve_document(pkg, transport=t, api_key=SECRET_KEY,
                           library_id="12345")
    assert rep.counts().get("create-failed") == 3
    assert any("Invalid item type" in m for m in rep.by_ref(1).messages)


# ---------------------------------------------------------------------------
# CSL -> Zotero JSON conversion unit tests
# ---------------------------------------------------------------------------

def test_csl_to_zotero_journal_article():
    z, notes = csl_to_zotero_item({
        "type": "article-journal",
        "title": "AI and labour",
        "author": [{"family": "Havryk", "given": "A"},
                   {"family": "Nazarova", "given": "T"}],
        "issued": {"date-parts": [["2024", 3, 3]]},
        "container-title": "Journal of Economics",
        "volume": "12", "issue": "3", "page": "1-20",
        "DOI": "10.46299/j.isjmef.20240303.01",
        "PMID": "38462539",
        "URL": "https://example.org",
    })
    assert z["itemType"] == "journalArticle"
    assert z["creators"][0] == {"creatorType": "author", "lastName": "Havryk",
                                "firstName": "A"}
    assert z["date"] == "2024-03-03"
    assert z["publicationTitle"] == "Journal of Economics"
    assert z["volume"] == "12" and z["issue"] == "3" and z["pages"] == "1-20"
    assert z["DOI"] == "10.46299/j.isjmef.20240303.01"
    assert z["extra"] == "PMID: 38462539"
    assert z["url"] == "https://example.org"
    assert notes == []


def test_csl_to_zotero_webpage_report_thesis():
    z, notes = csl_to_zotero_item({
        "type": "webpage", "title": "Revolut page",
        "author": [{"name": "Revolut"}],
        "issued": {"date-parts": [["2024", 5, 7]]},
        "publisher": "Revolut Ltd",           # no Zotero field on webpage
        "publisher-place": "London",
        "accessed": {"date-parts": [["2024", 6]]},
        "URL": "https://revolut.com",
    })
    assert z["itemType"] == "webpage"
    assert z["creators"] == [{"creatorType": "author", "name": "Revolut"}]
    assert z["date"] == "2024-05-07"
    assert z["accessDate"] == "2024-06"
    assert z["url"] == "https://revolut.com"
    assert "publisher" not in z and "place" not in z
    assert any("publisher not transferred" in n for n in notes)

    z, _ = csl_to_zotero_item({
        "type": "report", "title": "R", "number": "NIST AI 100-1",
        "publisher": "NIST", "publisher-place": "Gaithersburg",
        "issued": {"date-parts": [["2023"]]}},
    )
    assert z["reportNumber"] == "NIST AI 100-1"
    assert z["institution"] == "NIST" and z["place"] == "Gaithersburg"

    z, _ = csl_to_zotero_item({
        "type": "thesis", "title": "T", "genre": "PhD thesis",
        "publisher": "Some University",
        "issued": {"date-parts": [["2019"]]}},
    )
    assert z["thesisType"] == "PhD thesis"
    assert z["university"] == "Some University"


def test_csl_to_zotero_partial_dates_and_unknown_fields():
    z, notes = csl_to_zotero_item({
        "type": "article-journal", "title": "T",
        "issued": {"date-parts": [["2021", 7]]},
        "medium": "print",
    })
    assert z["date"] == "2021-07"
    assert any("medium" in n for n in notes)


# ---------------------------------------------------------------------------
# cache unit tests
# ---------------------------------------------------------------------------

def test_http_cache_roundtrip_and_disabled(tmp_path):
    db = tmp_path / "sub" / "cache.sqlite3"
    with HttpCache(db) as c:
        assert c.get("k", "missing") is None
        c.set("k", "q1", {"a": 1})
        assert c.get("k", "q1") == {"a": 1}
        c.set("k", "q1", {"a": 2})               # replace
        assert c.get("k", "q1") == {"a": 2}
    with HttpCache(db) as c2:                    # persists
        assert c2.get("k", "q1") == {"a": 2}
    off = HttpCache(None)
    assert off.get("k", "q1") is None
    off.set("k", "q1", {"x": 1})                 # no-op, no crash
    assert off.get("k", "q1") is None


# ---------------------------------------------------------------------------
# end-to-end on the real document (mocked HTTP)
# ---------------------------------------------------------------------------

def test_real_vancouver_resolution(real_docx):
    lib = [
        # the one item that already lives in the sample library
        lib_item("NU469JIQ", type="article-journal",
                 title="Artificial intelligence and its role in the labor market",
                 DOI="10.46299/j.isjmef.20240303.01",
                 author=[{"family": "Havryk", "given": "A"},
                         {"family": "Nazarova", "given": "T"}],
                 issued={"date-parts": [["2024", 6, 1]]}),
    ]
    t = FakeTransport(local_routes(lib)
                      + [crossref_route([]), pubmed_route({})])
    pkg = DocxPackage.open(real_docx)
    rep = resolve_document(pkg, transport=t)
    counts = rep.counts()
    # ref 4 is reused from the existing Zotero field (mapping FIXTURE-xO3paPuA)
    r4 = rep.by_ref(4)
    assert r4.status == "reused-existing-field"
    assert r4.item.key == "NU469JIQ"
    # everything else matched nothing in this one-item library
    assert counts.get("create-blocked-no-writes") == 15
    assert rep.by_ref(12).parsed_csl["type"] == "book"


# ---------------------------------------------------------------------------
# live test (skipped by default; enable with --run-live)
# ---------------------------------------------------------------------------

@pytest.mark.live
def test_live_local_api_read():
    access = ZoteroAccess(RequestsTransport(), timeout=3.0)
    assert access.probe_local(), "local Zotero is not running on localhost:23119"
    access.probe_local_writes()          # empty-body POST: cannot create anything
    items = access.fetch_items_csl()
    assert isinstance(items, list)
    print(access.mode_line())
